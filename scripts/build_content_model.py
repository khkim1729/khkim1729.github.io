#!/usr/bin/env python3
"""Build the Google-Sheets-ready portfolio workbook and offline JSON snapshot."""

from __future__ import annotations

import argparse
import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


PUBLICATION_HEADERS = [
    "Pub_ID", "Year", "Title", "Venue_Name", "Authors", "Paper_Link",
    "Venue_Link", "Notes", "Code", "Poster_Link", "Slides_link", "Cite",
    "In Google Scholar", "Cited at Least Once", "Status", "Type",
    "Oral", "Featured", "Display_Order", "Title_KO", "Venue_Name_KO",
    "Notes_KO", "Is_Visible", "Source",
]

CV_CONTENT_HEADERS = [
    "content_id", "section_id", "section_title_en", "section_title_ko",
    "entry_type", "title_en", "title_ko", "subtitle_en", "subtitle_ko",
    "period_en", "period_ko", "description_en", "description_ko", "link_url",
    "icon", "display_order", "is_visible",
]


EXTRA_PAPERS = [
    ("KHK_UR_001", 2026, "Automated dental caries segmentation in panoramic radiographs using a lightweight dual-stage deep learning framework", "Scientific Reports", "Yeongseok Seo, Kyeonghun Kim, Jong-yeol Lee, Jihun Kim, Dohyun Chun†"),
    ("KHK_UR_002", 2026, "Task-Side Geometry Shapes Terminal Class Geometry", "NeurIPS 2026", "Minseo Choi, Youngung Han, Kyeonghun Kim, Yoohyun Kim, Jewoo Park, Seungwon Park, Seoyoung Ju, Ken Ying-Kai Liao, Hyuk-Jae Lee, Nam-Joon Kim†"),
    ("KHK_UR_003", 2026, "Auditing Correlated Failures in Frozen-Feature Pretrained-Encoder Pools for Medical Segmentation", "NeurIPS 2026", "Eunseob Choi, Kyeonghun Kim, Hyuk-Jae Lee, Nam-Joon Kim†"),
    ("KHK_UR_004", 2026, "Auditing Capsule Vision 2024: Within-Split Train-to-Validation Re-Exposure and a Kvasir-Channel Sensitivity Diagnostic", "NeurIPS 2026", "Eunseob Choi, Kyeonghun Kim, Hyuk-Jae Lee, Nam-Joon Kim†"),
    ("KHK_UR_005", 2027, "FBI: A Framework for Benchmark Integrity via Keyed Option-Order Fingerprinting", "AAAI 2027", "Ina Jung, Kyeonghun Kim, Youngung Han, Seungwoo Baek, Jiwon Park, Minjeong Kim, Minseo Kim, Soo Yong Kim, Je Won Yeom, Sieun Hyeon†"),
    ("KHK_UR_006", 2027, "QDDPM: Quaternion Denoising Diffusion Probabilistic Models for Zero-Cost Super-Resolution", "AAAI 2027", "Seongheon Choi, Kyeonghun Kim, Insung Hwang, Taeyun Kim, Donghyeon Seo, Suemin Yang, Yoonseon Jung, Seoyoon Koo, Wonhyuk Kim, Pa Hong, Ken Ying-Kai Liao, Hyuk-Jae Lee, Nam-Joon Kim†"),
    ("KHK_UR_007", 2027, "Period Deserts: Controlled Corpus Interventions Induce and Repair Forecasting Failures in Time-Series Models", "AAAI 2027", "Eunseob Choi, Kyeonghun Kim, Hyuk-Jae Lee, Nam-Joon Kim†"),
    ("KHK_UR_008", 2026, "Hierarchy Aware Preference Optimization for the Safety of Korean Small Language Models", "EMNLP 2026", "Soo Yong Kim, Junyoung Koh, Kyeonghun Kim, Seunghyeok Hong†"),
    ("KHK_UR_009", 2027, "SSDenoDet: Self-Supervised Cross-Domain Feature Denoising for SAR Target Detection", "AAAI 2027", "Kyeonghun Kim, Yului Jeong, Youngung Han, Yunjin Seo, Youngseo Kim, Yeonghyeon Park, SeoJeong Woo, Giseong Hwang, Ken Ying-Kai Liao, Nam-Joon Kim†"),
    ("KHK_UR_010", 2027, "HERO: Hybrid PET-CT Representation Learning for Head and Neck Tumor Segmentation and Outcome Prediction", "AAAI 2027", "Kyeonghun Kim, Minju Kim, Tae Gyun Kim, Juyeon Yoo, Yehee Kim, Seohyeon Ji, Seokjun Choi, Kyung Seok Yuh, Pa Hong, Ken Ying-Kai Liao, Nam-Joon Kim†"),
]


PROJECTS = [
    ("PROJ_3DLLDM", "3d-lldm", "3D-LLDM", "3D-LLDM", "Label-guided 3D latent diffusion for high-resolution synthetic liver MRI.", "고해상도 합성 간 MRI를 위한 라벨 유도 3D 잠재 확산 모델입니다.", "Medical Imaging; Diffusion; MRI", "https://opengraph.githubassets.com/portfolio/khkim1729/3D-LLDM", "https://github.com/khkim1729/3D-LLDM", "https://arxiv.org/abs/2603.23845"),
    ("PROJ_TIME", "time", "TIME", "TIME", "A 2.5D IDH-predictive multimodal ensemble for multitask glioma subtyping.", "다중 과제 교종 아형 분류를 위한 2.5D IDH 예측 멀티모달 앙상블입니다.", "Medical Imaging; Glioma; Multimodal", "https://opengraph.githubassets.com/portfolio/khkim1729/TIME", "https://github.com/khkim1729/TIME", ""),
    ("PROJ_MATHENA", "mathena", "MATHENA", "MATHENA", "A hierarchical tooth anatomy estimator built with Mamba-based representations.", "Mamba 표현을 활용한 계층적 치아 해부학 추정 모델입니다.", "Dental AI; Mamba; Segmentation", "https://opengraph.githubassets.com/portfolio/khkim1729/mathena", "https://github.com/khkim1729/mathena", "https://arxiv.org/abs/2604.00537"),
    ("PROJ_MAESIL", "maesil", "MAESIL / MAEgic", "MAESIL / MAEgic", "Enhanced self-supervised medical image learning with masked autoencoders.", "마스킹 오토인코더 기반 의료영상 자기지도학습 프로젝트입니다.", "Self-supervised Learning; Medical Imaging", "https://opengraph.githubassets.com/portfolio/khkim1729/MAEgic", "https://github.com/khkim1729/MAEgic", "https://arxiv.org/abs/2604.00514"),
    ("PROJ_COTTA", "cotta", "COTTA", "COTTA", "Context-aware transfer adaptation for autonomous-driving trajectory prediction.", "자율주행 궤적 예측을 위한 문맥 인식 전이 적응 모델입니다.", "Autonomous Driving; Trajectory Prediction", "https://opengraph.githubassets.com/portfolio/khkim1729/COTTUS", "https://github.com/khkim1729/COTTUS", "https://arxiv.org/abs/2604.00402"),
    ("PROJ_LUCAS", "lucas", "LUCAS", "LUCAS", "Deep-learning analysis of multi-phase contrast-enhanced ultrasound.", "다중 시기 조영증강 초음파 딥러닝 분석 프로젝트입니다.", "Medical Imaging; CEUS; Ultrasound", "https://opengraph.githubassets.com/portfolio/khkim1729/LUCAS", "https://github.com/khkim1729/LUCAS", ""),
    ("PROJ_CIPHER", "cipher", "CIPHER", "CIPHER", "Representation learning for detecting GAN- and diffusion-generated counterfeit images.", "GAN·확산 모델 생성 위조 이미지 탐지를 위한 표현학습 프로젝트입니다.", "Generative AI; Detection; Representation Learning", "https://opengraph.githubassets.com/portfolio/outta-ai/2025_OUTTA_AIBootcamp_Deep-Learning_final_project", "https://github.com/outta-ai/2025_OUTTA_AIBootcamp_Deep-Learning_final_project", "https://arxiv.org/abs/2603.29356"),
    ("PROJ_PORTFOLIO", "portfolio", "Kyeonghun Kim Portfolio", "김경훈 포트폴리오", "A bilingual, spreadsheet-driven academic portfolio and publication search experience.", "구글시트 기반으로 관리되는 다국어 연구 포트폴리오와 논문 검색 환경입니다.", "Web; Portfolio; Google Sheets", "https://opengraph.githubassets.com/portfolio/khkim1729/khkim1729.github.io", "https://github.com/khkim1729/khkim1729.github.io", "https://docs.google.com/presentation/d/1Ns0ebR9yU-mzKa_Ua9e0_YRzSjMfQzVUia90qfQuW5E/edit?usp=sharing"),
]


def clean(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def table(ws):
    rows = [[clean(v) for v in row] for row in ws.iter_rows(values_only=True)]
    headers = [str(v) for v in rows[0]]
    return [dict(zip(headers, row)) for row in rows[1:] if any(v != "" for v in row)]


def status_for(venue):
    return "Under Review" if "under review" in venue.lower() else "Published"


def make_publications(imsi):
    source = table(imsi["WEB_Publications"])
    selected = []
    for order, row in enumerate(reversed(source), 1):
        if "Kyeonghun Kim" not in str(row.get("Authors", "")):
            continue
        venue = str(row.get("Venue_Name", ""))
        selected.append([
            *[row.get(h, "") for h in PUBLICATION_HEADERS[:14]],
            status_for(venue),
            "Journal" if any(x in venue.lower() for x in ("ultrasonics", "scientific reports")) else "Conference",
            "oral" in venue.lower(),
            order <= 6,
            order * 10,
            "", "", "", True, "IMSI WEB_Publications",
        ])
    base_order = len(selected) + 1
    for offset, (pid, year, title, venue, authors) in enumerate(EXTRA_PAPERS):
        selected.insert(0, [
            pid, year, title, venue, authors, "", "", "Under review", "", "", "", "",
            "-", "-", "Under Review", "Journal" if venue == "Scientific Reports" else "Conference",
            False, offset < 4, (base_order + offset) * 10, "", "", "심사 중", True, "Personal",
        ])
    return selected


def people_rows(imsi, publications):
    people = table(imsi["DB_People"])
    by_name = {str(p.get("Name_EN_FULL", "")).strip(): p for p in people}
    names = []
    for row in publications:
        for raw in str(row[4]).split(","):
            name = re.sub(r"[\*†‡]+$", "", raw.strip())
            if name and name not in names:
                names.append(name)
    result = []
    for idx, name in enumerate(names, 1):
        person = by_name.get(name, {})
        url = person.get("Linkedin") or person.get("Github") or person.get("Blog") or ""
        if name == "Kyeonghun Kim":
            url = "https://khkim1729.github.io"
        result.append([
            person.get("rn") or f"EXT_{idx:04d}", person.get("Name_KR", ""), name,
            person.get("Github", ""), person.get("Blog", ""), person.get("Linkedin", ""),
            url, True,
        ])
    return result


def credential_rows(path):
    ws = load_workbook(path, data_only=True).active
    data = table(ws)
    rows = []
    for i, row in enumerate(data, 1):
        rows.append([
            f"CRED_{i:03d}", row.get("자격증명", ""), row.get("Certification / License", ""),
            row.get("등급", ""), row.get("Level / Grade", ""), row.get("취득일", ""),
            row.get("Date Acquired", ""), row.get("발급기관", ""),
            row.get("Issuing Organization", ""), row.get("발급 번호", ""),
            row.get("Credential ID", ""), i * 10, True,
        ])
    return rows


def add_sheet(wb, title, headers, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i in range(1, len(headers) + 1):
        width = widths.get(headers[i - 1], 18) if widths else 18
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width
    ws.sheet_view.showGridLines = False
    return ws


def build(args):
    imsi = load_workbook(args.imsi, data_only=True)
    publications = make_publications(imsi)
    people = people_rows(imsi, publications)
    credentials = credential_rows(args.credentials)
    cv_path = Path(__file__).resolve().parents[1] / "assets/data/cv-content.json"
    cv_content = json.loads(cv_path.read_text(encoding="utf-8"))

    wb = Workbook()
    wb.remove(wb.active)
    add_sheet(wb, "README", ["Item", "Guideline"], [
        ["Editing flow", "Google Sheets에 업로드한 뒤 DB_* 탭만 편집합니다. WEB_* 탭은 웹 공개용입니다."],
        ["IMSI copy/paste", "IMSI WEB_Publications의 A:N 열을 DB_Publications A:N에 그대로 붙여넣을 수 있습니다."],
        ["Languages", "한국어 열이 비어 있으면 Apps Script와 웹사이트가 영어 값을 fallback으로 사용합니다."],
        ["Projects/images", "Projects의 cover_image_url과 Project_Content의 media_url에 공개 HTTPS 이미지 URL을 입력합니다."],
        ["Deploy", "docs/GOOGLE_SHEETS_DEPLOYMENT.md와 docs/google-sheets-code.gs를 사용해 웹 앱으로 배포합니다."],
        ["Public contact", "API에는 khkim1729@gmail.com과 자격증 번호를 공개합니다. 개인 전화번호는 입력하지 않습니다."],
    ])
    add_sheet(wb, "Site_Config", ["config_key", "value_en", "value_ko", "data_type", "notes"], [
        ["site_title", "Kyeonghun Kim | AI Researcher", "김경훈 | AI 연구자", "text", "SEO/title"],
        ["site_description", "Medical imaging, generative AI, multimodal learning, and hyperspectral imaging.", "의료영상, 생성형 AI, 멀티모달 학습 및 초분광 영상 연구.", "text", "SEO description"],
        ["default_language", "en", "ko", "enum", "en/ko"],
        ["github_url", "https://github.com/khkim1729", "https://github.com/khkim1729", "url", ""],
        ["google_scholar_url", "https://scholar.google.com/citations?user=N3LVcyEAAAAJ&hl=en", "https://scholar.google.com/citations?user=N3LVcyEAAAAJ&hl=ko", "url", ""],
        ["portfolio_url", "https://docs.google.com/presentation/d/1Ns0ebR9yU-mzKa_Ua9e0_YRzSjMfQzVUia90qfQuW5E/edit?usp=sharing", "https://docs.google.com/presentation/d/1Ns0ebR9yU-mzKa_Ua9e0_YRzSjMfQzVUia90qfQuW5E/edit?usp=sharing", "url", ""],
        ["public_email", "khkim1729@gmail.com", "khkim1729@gmail.com", "email", "Public contact email; do not add a phone number."],
    ])
    add_sheet(wb, "Home_Sections", ["section_id", "title_en", "title_ko", "body_en", "body_ko", "display_order", "is_visible"], [
        ["ABOUT", "About Me", "소개", "AI researcher working across medical imaging, generative AI, multimodal learning, and hyperspectral imaging. Research Assistant at IMSI Lab and AI Researcher at GNEWSOFT.", "의료영상, 생성형 AI, 멀티모달 학습, 초분광 영상을 연구하는 AI 연구자입니다. IMSI Lab 연구조교이자 GNEWSOFT AI 연구원으로 활동하고 있습니다.", 10, True],
        ["RESEARCH", "Research & Leadership", "연구 및 리더십", "I lead medical imaging research teams, maintain the IMSI Lab website, and collaborate with clinical and NVIDIA researchers.", "의료영상 연구팀을 이끌고 IMSI Lab 웹사이트를 운영하며 임상·NVIDIA 연구진과 협업하고 있습니다.", 20, True],
    ])
    add_sheet(
        wb,
        "CV_Content",
        CV_CONTENT_HEADERS,
        [[row.get(header, "") for header in CV_CONTENT_HEADERS] for row in cv_content],
        {"section_title_en": 28, "section_title_ko": 24, "title_en": 42,
         "title_ko": 35, "description_en": 75, "description_ko": 75,
         "link_url": 50},
    )
    add_sheet(wb, "Experience_Education", ["entry_id", "category", "organization_en", "organization_ko", "role_en", "role_ko", "period_en", "period_ko", "description_en", "description_ko", "display_order", "is_visible"], [
        ["EXP_001", "Research", "IMSI Lab, Seoul National University", "서울대학교 IMSI Lab", "Research Assistant", "연구조교", "Sep. 2024 – Present", "2024년 9월 – 현재", "Multimodal medical AI research and team leadership.", "멀티모달 의료 AI 연구 및 팀 리더십.", 10, True],
        ["EXP_002", "Professional", "GNEWSOFT", "지뉴스프트", "AI Researcher", "AI 연구원", "Oct. 2024 – Present", "2024년 10월 – 현재", "Foundation models and applied AI systems.", "파운데이션 모델 및 응용 AI 시스템 개발.", 20, True],
        ["EDU_001", "Education", "Kookmin University", "국민대학교", "M.S. in Artificial Intelligence", "인공지능학과 석사", "Sep. 2022 – Feb. 2024", "2022년 9월 – 2024년 2월", "GPA 4.03/4.5", "학점 4.03/4.5", 30, True],
        ["EDU_002", "Education", "Hankyong National University", "한경국립대학교", "B.S. in Computer Science", "컴퓨터공학과 학사", "Mar. 2018 – Aug. 2022", "2018년 3월 – 2022년 8월", "GPA 4.12/4.5; 158 credits.", "학점 4.12/4.5; 158학점 이수.", 40, True],
    ])
    add_sheet(wb, "DB_People", ["Person_ID", "Name_KO", "Name_EN", "Github", "Blog", "Linkedin", "Primary_URL", "Is_Visible"], people, {"Name_KO": 18, "Name_EN": 24, "Primary_URL": 45})
    pubs_ws = add_sheet(wb, "DB_Publications", PUBLICATION_HEADERS, publications, {"Title": 55, "Venue_Name": 42, "Authors": 70, "Paper_Link": 42, "Cite": 55})
    status_validation = DataValidation(type="list", formula1='"Published,Accepted,Under Review,Preprint"', allow_blank=False)
    pubs_ws.add_data_validation(status_validation)
    status_validation.add(f"O2:O{max(2, pubs_ws.max_row)}")
    author_rows = []
    person_ids = {row[2]: row[0] for row in people}
    for pub in publications:
        for order, raw in enumerate(str(pub[4]).split(","), 1):
            marked = raw.strip()
            name = re.sub(r"[\*†‡]+$", "", marked)
            author_rows.append([pub[0], order, person_ids.get(name, ""), name, "*" in marked, "†" in marked])
    add_sheet(wb, "DB_PublicationAuthor", ["Pub_ID", "Author_Order", "Person_ID", "Author_Name_Display", "Is_CoFirst_Author", "Is_Corresponding"], author_rows)
    # Keep a value-only compatibility view. ARRAYFORMULA is not evaluated by
    # Excel/openpyxl and can become #NAME? when the XLSX is first uploaded.
    add_sheet(wb, "WEB_Publications", PUBLICATION_HEADERS, publications, {"Title": 55, "Venue_Name": 42, "Authors": 70})
    add_sheet(wb, "WEB_People", ["Person_ID", "Name_KO", "Name_EN", "Primary_URL"], [[p[0], p[1], p[2], p[6]] for p in people])
    add_sheet(wb, "Projects", ["project_id", "slug", "title_en", "title_ko", "summary_en", "summary_ko", "tags", "cover_image_url", "github_link", "paper_or_demo_link", "featured_on_home", "display_order", "is_visible"], [
        [*p, i < 4, i * 10, True] for i, p in enumerate(PROJECTS, 1)
    ], {"summary_en": 55, "summary_ko": 55, "cover_image_url": 55, "github_link": 45})
    add_sheet(wb, "Project_Content", ["content_id", "project_id", "block_type", "title_en", "title_ko", "content_en", "content_ko", "media_url", "caption_en", "caption_ko", "display_order", "is_visible"], [
        [f"PC_{i:03d}", p[0], "paragraph", "Overview", "개요", p[4], p[5], p[7], p[2], p[3], 10, True] for i, p in enumerate(PROJECTS, 1)
    ])
    add_sheet(wb, "Professional_Credentials", ["credential_id", "name_ko", "name_en", "level_ko", "level_en", "date_ko", "date_en", "issuer_ko", "issuer_en", "credential_number", "verification_id", "display_order", "is_visible"], credentials)
    add_sheet(wb, "News", ["news_id", "date", "title_en", "title_ko", "description_en", "description_ko", "link_url", "featured_on_home", "display_order", "is_visible"], [
        ["NEWS_001", "2026-07-27", "Portfolio content moved to Google Sheets", "포트폴리오 콘텐츠를 Google Sheets로 이전", "Publications, projects, and credentials can now be maintained without code changes.", "논문·프로젝트·자격 정보를 코드 수정 없이 관리할 수 있습니다.", "", True, 10, True],
    ])
    add_sheet(wb, "Skills", ["skill_id", "category_en", "category_ko", "skill_name", "display_order", "is_visible"], [
        ["SK_001", "Programming", "프로그래밍", "Python", 10, True],
        ["SK_002", "AI/ML", "AI/ML", "PyTorch / MONAI / Transformers", 20, True],
        ["SK_003", "Domains", "전문 분야", "Medical Imaging / Hyperspectral Imaging", 30, True],
    ])
    add_sheet(wb, "Developer_Notes", ["Topic", "Recommendation"], [
        ["Public API", "공개 이메일과 credential_number/verification_id를 반환합니다. 개인 전화번호 열은 만들거나 입력하지 않습니다."],
        ["Fallback", "GitHub Pages는 assets/data/portfolio-data.json을 오프라인 fallback으로 사용합니다."],
        ["lang=all", "모든 언어 열을 반환합니다. lang=en/ko는 *_en, *_ko 열을 공통 키로 정규화합니다."],
    ])

    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)

    sheets = {
        "Site_Config": table(wb["Site_Config"]),
        "Home_Sections": table(wb["Home_Sections"]),
        "CV_Content": table(wb["CV_Content"]),
        "DB_People": table(wb["DB_People"]),
        "Publications": [dict(zip(PUBLICATION_HEADERS, row)) for row in publications],
        "Projects": table(wb["Projects"]),
        "Project_Content": table(wb["Project_Content"]),
        "Professional_Credentials": table(wb["Professional_Credentials"]),
        "News": table(wb["News"]),
        "Skills": table(wb["Skills"]),
    }
    args.json_output.parent.mkdir(parents=True, exist_ok=True)
    args.json_output.write_text(json.dumps({"generated_at": datetime.now().isoformat(), "sheets": sheets}, ensure_ascii=False, indent=2), encoding="utf-8")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--imsi", type=Path, required=True)
    parser.add_argument("--credentials", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--json-output", type=Path, required=True)
    build(parser.parse_args())


if __name__ == "__main__":
    main()
